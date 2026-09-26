"""종합 레포트 양식을 계산하고, rules.yaml 기준으로 섹션을 고릅니다."""

from __future__ import annotations

from datetime import datetime, timedelta, timezone
from pathlib import Path
from xml.sax.saxutils import escape

import openpyxl
import yaml

from reportkit.empty import as_text, is_empty
from reportkit.errors import ReportError
from reportkit.office import recalculate

ROOT = Path(__file__).resolve().parent
RULES_PATH = ROOT / "rules.yaml"
LOGIC_PATH = ROOT / "logic" / "recovery_template.xlsx"
KST = timezone(timedelta(hours=9))


def load_rules(path: Path | None = None) -> dict:
    path = Path(path or RULES_PATH)
    data = yaml.safe_load(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict) or not data.get("sections"):
        raise ReportError("종합 레포트 규칙 파일에 섹션이 없습니다.")
    return data


def build_model(workbook_path: Path, rules: dict | None = None) -> dict:
    rules = rules or load_rules()
    tokens = {str(item).strip().casefold() for item in rules.get("empty_tokens") or []}
    missing_text = rules["document"]["missing_text"]
    loc = rules["locators"]

    calculated = Path(workbook_path).with_suffix(".calculated.xlsx")
    recalculate(workbook_path, calculated)
    source = openpyxl.load_workbook(workbook_path, data_only=False)
    book = openpyxl.load_workbook(calculated, data_only=True)
    _require(book, list(loc.values())[:7] if False else [
        loc["cover_sheet"], loc["summary_sheet"], loc["session_sheet"],
        loc["capability_sheet"], loc["functional_sheet"], loc["cancer_sheet"], loc["next_sheet"],
    ])

    filled, total, missing_labels = _yellow_inputs(source, tokens)
    content = _read_content(book, loc, tokens, missing_text)
    presence = _presence(content)

    sections = []
    omitted = []
    for spec in rules["sections"]:
        core = list(spec.get("core") or [])
        missing_core = [key for key in core if not presence.get(key)]
        policy = spec.get("core_policy") or "any_missing"
        if policy == "all_missing":
            drop = bool(core) and len(missing_core) == len(core)
        else:
            drop = bool(missing_core)
        if drop:
            omitted.append({
                "id": spec["id"],
                "title": spec["title"],
                "reason": "핵심 항목 없음: " + ", ".join(missing_core),
            })
            continue
        partial = bool(missing_core)
        block = _render_section(spec, content, missing_text, partial, spec.get("on_partial") or "show_placeholder")
        block["status"] = "partial" if partial else "full"
        block["reason"] = (
            "일부 항목만 있어 있는 내용만 표시했습니다."
            if partial and spec.get("on_partial") == "render_available"
            else "일부 항목만 있어 빈 항목은 없음 문구로 표시했습니다."
            if partial
            else "필요 항목이 있어 전체를 표시했습니다."
        )
        sections.append(block)

    return {
        "document": rules["document"],
        "version": str(rules["version"]),
        "updated_at": str(rules["updated_at"]),
        "content": content,
        "sections": sections,
        "omitted": omitted,
        "filled_count": filled,
        "total_count": total,
        "missing_labels": missing_labels,
        "generated_at": datetime.now(KST).strftime("%Y-%m-%d %H:%M"),
        "source_name": Path(workbook_path).name,
        "tag_labels": rules.get("tag_labels") or {},
        "note_labels": rules.get("note_labels") or {},
        "practice_groups": _practice_groups(content["practice_rows"]),
    }


def _require(book, names: list[str]) -> None:
    missing = [name for name in names if name not in book.sheetnames]
    if missing:
        raise ReportError("종합 레포트 양식이 아닙니다. 시트가 없습니다: " + ", ".join(missing))


def _yellow_inputs(book, tokens: set[str]) -> tuple[int, int, list[str]]:
    filled = 0
    total = 0
    missing = []
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if not _is_yellow(cell):
                    continue
                total += 1
                if is_empty(cell.value, tokens):
                    missing.append(f"{sheet.title} {cell.coordinate}")
                else:
                    filled += 1
    return filled, total, missing


def _is_yellow(cell) -> bool:
    fill = cell.fill
    if fill is None or fill.patternType != "solid" or fill.fgColor is None:
        return False
    rgb = fill.fgColor.rgb
    return isinstance(rgb, str) and rgb.upper().endswith("FFF7CC")


def _read_content(book, loc: dict, tokens: set[str], missing_text: str) -> dict:
    cover = book[loc["cover_sheet"]]
    summary = book[loc["summary_sheet"]]
    sessions = book[loc["session_sheet"]]
    capability = book[loc["capability_sheet"]]
    functional = book[loc["functional_sheet"]]
    cancer = book[loc["cancer_sheet"]]
    nxt = book[loc["next_sheet"]]

    def key_value(sheet, key: str, column: int = 3) -> str:
        for row in range(1, sheet.max_row + 1):
            if as_text(sheet.cell(row, 1).value) == key:
                return as_text(sheet.cell(row, column).value)
        return ""

    visits = _table(summary, loc["visit_header"], ["no", "date", "status"])
    session_rows = _table(sessions, loc["visit_header"], ["no", "date", "tag", "title", "desc", "note_type", "note_text"])
    capabilities = _table(capability, loc["capability_header"], ["name", "before", "after", "note"])
    learned = _table(functional, loc["learned_header"], ["icon", "label", "range"])
    sppb = _table(functional, loc["sppb_header"], ["label", "caption", "unit", "first", "last", "note", "score_first", "score_last"])
    arm_rows = _arm_rows(cancer, loc["arm_header"], tokens)
    rom_rows = _table(cancer, loc["rom_header"], ["motion", "before", "after", "unit", "desc", "delta"])

    return {
        "name": key_value(cover, "name"),
        "cancer_type": key_value(cover, "cancerType"),
        "date_range": key_value(cover, "dateRange"),
        "visit_count": key_value(cover, "총 방문 횟수"),
        "issued_on": key_value(cover, "리포트 발행일"),
        "visits": visits,
        "summary_stats": _stats(summary, "값(value)", 3),
        "summary_narrative": key_value(summary, "summary.narrative"),
        "practice_rows": _grouped(summary, "그룹명(group)"),
        "practice_narrative": key_value(summary, "practice.narrative"),
        "care_stats": _pairs_after(summary, "값(value)", second_header="라벨(label)", stop="care.narrative"),
        "care_narrative": key_value(summary, "care.narrative"),
        "call_count": key_value(cancer if False else summary, "careDetail.callCount"),
        "professor": key_value(summary, "careDetail.professorFeedback"),
        "next_plan": key_value(summary, "careDetail.nextPlanText"),
        "next_chips": key_value(summary, "careDetail.nextPlanChips"),
        "sessions": session_rows,
        "session_note": key_value(sessions, "sessionLog.footnote"),
        "capabilities": capabilities,
        "capability_quote": key_value(capability, "capabilityChanges.closingQuote"),
        "learned": learned,
        "learned_badge": key_value(functional, "learnedExercises.badge"),
        "sppb": sppb,
        "sppb_badge": key_value(functional, "sppb.badge"),
        "sppb_total": key_value(functional, "sppb.total"),
        "sppb_headline": key_value(functional, "sppb.interpHeadline"),
        "sppb_detail": key_value(functional, "sppb.interpDetail"),
        "weight": _weight(functional),
        "weight_note": key_value(functional, "weight.note"),
        "bmi_text": key_value(functional, "weight.bmi (첫/재)"),
        "cancer_type_metric": key_value(cancer, "cancerSpecific.type"),
        "arm_rows": arm_rows,
        "arm_svg": _arm_svg(arm_rows),
        "arm_badge": key_value(cancer, "armCircumference.badge"),
        "arm_delta": key_value(cancer, "armCircumference.deltaLabel"),
        "arm_interp": key_value(cancer, "armCircumference.interp"),
        "volume": _metric_row(cancer, "표준화 5지점 부피 (volume)"),
        "relative": _metric_row(cancer, "양팔 상대 부피 차이 (relativeDiff)"),
        "arm_footnote": key_value(cancer, "armCircumference.footnote"),
        "rom_rows": rom_rows,
        "rom_headline": key_value(cancer, "shoulderROM.interpHeadline"),
        "rom_detail": key_value(cancer, "shoulderROM.interpDetail"),
        "rom_footnote": key_value(cancer, "shoulderROM.footnote"),
        "other_cancer_note": key_value(cancer, "cancerSpecific.otherCancerNote"),
        "next_date": key_value(nxt, "nextSteps.dateLabel"),
        "next_keep": key_value(nxt, "nextSteps.keep"),
        "next_watch": key_value(nxt, "nextSteps.watch"),
        "next_ask": key_value(nxt, "nextSteps.askDoctor"),
        "nurse_author": key_value(nxt, "nextSteps.nurseNote.author"),
        "nurse_text": key_value(nxt, "nextSteps.nurseNote.text"),
        "closing": key_value(nxt, "nextSteps.closingNote"),
        "missing_text": missing_text,
    }


def _presence(content: dict) -> dict[str, bool]:
    def rows_have(rows, key):
        return any(not is_empty(row.get(key)) for row in rows)

    return {
        "name": not is_empty(content["name"]),
        "visit_no": rows_have(content["visits"], "no"),
        "session_title": rows_have(content["sessions"], "title"),
        "capability_name": rows_have(content["capabilities"], "name"),
        "learned_label": rows_have(content["learned"], "label"),
        "sppb_label": rows_have(content["sppb"], "label"),
        "weight_first": not is_empty((content["weight"] or {}).get("first")),
        "arm_value": rows_have(content["arm_rows"], "value"),
        "rom_motion": rows_have(content["rom_rows"], "motion"),
        "next_keep": not is_empty(content["next_keep"]),
        "next_watch": not is_empty(content["next_watch"]),
        "next_ask": not is_empty(content["next_ask"]),
    }


def _find_header(sheet, header: str) -> int | None:
    for row in range(1, sheet.max_row + 1):
        if as_text(sheet.cell(row, 1).value) == header:
            return row
    return None


def _table(sheet, header: str, keys: list[str]) -> list[dict]:
    start = _find_header(sheet, header)
    if start is None:
        return []
    rows = []
    for row in range(start + 1, sheet.max_row + 1):
        values = [sheet.cell(row, col).value for col in range(1, len(keys) + 1)]
        if all(is_empty(value) for value in values):
            break
        if isinstance(values[0], str) and _is_section_break(str(values[0])):
            break
        item = {key: as_text(value) for key, value in zip(keys, values)}
        if any(item.values()):
            rows.append(item)
    return rows


def _is_section_break(text: str) -> bool:
    if text.startswith(("①", "②", "③", "※")):
        return True
    if "." in text and " " not in text.split(".")[0]:
        return True
    return False


def _stats(sheet, header: str, width: int) -> list[dict]:
    return [
        {"value": row.get("no", ""), "label": row.get("date", ""), "sub": row.get("status", "")}
        for row in _table(sheet, header, ["no", "date", "status"])
        if row.get("date") and not str(row.get("date")).startswith("라벨")
    ] if False else _labeled_stats(sheet)


def _labeled_stats(sheet) -> list[dict]:
    header = _find_header(sheet, "값(value)")
    if header is None:
        return []
    rows = []
    for row in range(header + 1, sheet.max_row + 1):
        marker = as_text(sheet.cell(row, 1).value)
        if marker in {"summary.narrative", "practice.narrative", "care.narrative"} or marker.startswith("②"):
            break
        value = as_text(sheet.cell(row, 1).value)
        label = as_text(sheet.cell(row, 2).value)
        sub = as_text(sheet.cell(row, 3).value)
        if not value and not label:
            continue
        if label == "라벨(label)":
            continue
        rows.append({"value": value, "label": label, "sub": sub})
    return rows


def _grouped(sheet, header: str) -> list[dict]:
    start = _find_header(sheet, header)
    if start is None:
        return []
    rows = []
    for row in range(start + 1, sheet.max_row + 1):
        group = as_text(sheet.cell(row, 1).value)
        if not group or group.endswith(".narrative") or group.startswith("③"):
            break
        rows.append({
            "group": group,
            "value": as_text(sheet.cell(row, 2).value),
            "label": as_text(sheet.cell(row, 3).value),
        })
    return rows


def _pairs_after(sheet, header: str, second_header: str, stop: str) -> list[dict]:
    # care stats use the second '값(value)' header. Find the one followed by care rows.
    matches = []
    for row in range(1, sheet.max_row + 1):
        if as_text(sheet.cell(row, 1).value) == header and as_text(sheet.cell(row, 2).value) == second_header:
            if as_text(sheet.cell(row, 3).value) == "":
                matches.append(row)
    if not matches:
        return []
    start = matches[-1]
    rows = []
    for row in range(start + 1, sheet.max_row + 1):
        marker = as_text(sheet.cell(row, 1).value)
        if marker == stop or marker.startswith("담당자"):
            break
        if not marker:
            continue
        rows.append({"value": marker, "label": as_text(sheet.cell(row, 2).value)})
    return rows


def _weight(sheet) -> dict:
    for row in range(1, sheet.max_row + 1):
        if as_text(sheet.cell(row, 1).value) == "체중":
            return {
                "first": as_text(sheet.cell(row, 2).value),
                "last": as_text(sheet.cell(row, 3).value),
                "unit": as_text(sheet.cell(row, 4).value),
                "delta": as_text(sheet.cell(row, 5).value),
            }
    return {}


def _metric_row(sheet, label: str) -> dict:
    for row in range(1, sheet.max_row + 1):
        if as_text(sheet.cell(row, 1).value) == label:
            return {
                "label": label,
                "first": as_text(sheet.cell(row, 2).value),
                "last": as_text(sheet.cell(row, 3).value),
                "unit": as_text(sheet.cell(row, 4).value),
                "delta": as_text(sheet.cell(row, 5).value),
            }
    return {}


def _arm_rows(sheet, header: str, tokens: set[str]) -> list[dict]:
    start = None
    for row in range(1, sheet.max_row + 1):
        label = as_text(sheet.cell(row, 1).value)
        unit = as_text(sheet.cell(row, 2).value)
        if label == header and ("팔둘레" in unit or "cm" in unit.lower()):
            start = row
            break
    if start is None:
        return []
    rows = []
    for row in range(start + 1, sheet.max_row + 1):
        label = as_text(sheet.cell(row, 1).value)
        if not label or _is_section_break(label) or label.startswith("arm") or label.startswith("지표"):
            break
        value = sheet.cell(row, 2).value
        if is_empty(value, tokens):
            continue
        rows.append({"label": label, "value": as_text(value)})
    return rows


def _practice_groups(rows: list[dict]) -> list[dict]:
    groups = []
    for row in rows:
        name = row.get("group") or ""
        if not groups or groups[-1]["name"] != name:
            groups.append({"name": name, "items": []})
        groups[-1]["items"].append(row)
    return groups


def _arm_svg(rows: list[dict]) -> str:
    numbers = []
    for row in rows:
        try:
            numbers.append(float(str(row.get("value") or "").replace(",", "")))
        except ValueError:
            return ""
    if len(numbers) < 2:
        return ""
    width, height = 640, 128
    low, high = min(numbers), max(numbers)
    if high <= low:
        high = low + 1
    pad_l, pad_r, pad_t, pad_b = 28, 12, 18, 24
    inner_w = width - pad_l - pad_r
    inner_h = height - pad_t - pad_b

    def xpos(index: int) -> float:
        return pad_l + inner_w * index / (len(numbers) - 1)

    def ypos(value: float) -> float:
        return pad_t + inner_h * (high - value) / (high - low)

    points = " ".join(f"{xpos(index):.1f},{ypos(value):.1f}" for index, value in enumerate(numbers))
    marks = []
    for index, value in enumerate(numbers):
        label = escape(str(rows[index].get("label") or ""))
        shown = escape(str(rows[index].get("value") or ""))
        marks.append(f'<circle cx="{xpos(index):.1f}" cy="{ypos(value):.1f}" r="3.5" fill="#1d4f4c"/>')
        marks.append(
            f'<text x="{xpos(index):.1f}" y="{ypos(value) - 7:.1f}" text-anchor="middle" font-size="11" fill="#143840">{shown}</text>'
        )
        marks.append(
            f'<text x="{xpos(index):.1f}" y="{height - 6}" text-anchor="middle" font-size="11" fill="#5c6a66">{label}</text>'
        )
    return (
        f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {width} {height}" width="100%" font-family="ReportKR, Malgun Gothic, Apple SD Gothic Neo, sans-serif">'
        f'<polyline points="{points}" fill="none" stroke="#1d4f4c" stroke-width="2"/>'
        + "".join(marks)
        + "</svg>"
    )


def _render_section(spec, content, missing_text, partial: bool, mode: str) -> dict:
    show_blank = mode != "render_available"
    return {
        "id": spec["id"],
        "title": spec["title"],
        "show_blank": show_blank,
        "missing_text": missing_text,
        "partial": partial,
    }
