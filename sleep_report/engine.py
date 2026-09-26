"""로직 엑셀을 계산한 뒤, rules.yaml에 적힌 기준으로 섹션을 고릅니다."""

from __future__ import annotations

from datetime import datetime, timedelta, timezone
from pathlib import Path

import openpyxl
import yaml

from reportkit.empty import as_text, is_empty
from reportkit.errors import ReportError
from reportkit.office import recalculate
from sleep_report.charts import combo_svg, dual_svg, hours_label, line_svg, stacked_svg

ROOT = Path(__file__).resolve().parent
RULES_PATH = ROOT / "rules.yaml"
LOGIC_PATH = ROOT / "logic" / "sleep_logic.xlsx"
KST = timezone(timedelta(hours=9))


def load_rules(path: Path | None = None) -> dict:
    path = Path(path or RULES_PATH)
    data = yaml.safe_load(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict) or not data.get("sections"):
        raise ReportError("수면 로직 규칙 파일에 섹션이 없습니다.")
    return data


def build_model(workbook_path: Path, rules: dict | None = None) -> dict:
    rules = rules or load_rules()
    tokens = {str(item).strip().casefold() for item in rules.get("empty_tokens") or []}
    missing_text = rules["document"]["missing_text"]

    calculated = workbook_path.with_suffix(".calculated.xlsx")
    recalculate(workbook_path, calculated)
    book = openpyxl.load_workbook(calculated, data_only=True)
    _require_sheets(book, ["S1_Input_Raw", "S2_Input_Profile", "S4_Calc_Daily", "S7_Output_Summary"])

    raw = book["S1_Input_Raw"]
    profile = book["S2_Input_Profile"]
    daily = book["S4_Calc_Daily"]
    summary = book["S7_Output_Summary"]

    raw_headers = _headers(raw, 4)
    daily_headers = _headers(daily, 4)
    day_rows = _data_rows(raw)
    days = []
    for row in day_rows:
        label = as_text(raw.cell(row, 1).value) or f"{row}"
        raw_values = {name: raw.cell(row, col).value for name, col in raw_headers.items()}
        daily_values = {}
        daily_row = _matching_daily_row(daily, label, row)
        for name, col in daily_headers.items():
            daily_values[name] = daily.cell(daily_row, col).value
        days.append({"label": label, "row": row, "raw": raw_values, "daily": daily_values})

    monitor_headers = list(rules.get("monitor_headers") or [])
    monitored_flags = []
    for day in days:
        present = all(not is_empty(day["raw"].get(name), tokens) for name in monitor_headers)
        day["monitored"] = present
        monitored_flags.append(present)

    profile_fields = _profile_fields(profile, tokens)
    input_total, input_filled, missing_labels = _count_inputs(days, raw_headers, profile_fields, tokens)

    pages = _summary_pages(summary)
    weekly = book["S5_Calc_Weekly"] if "S5_Calc_Weekly" in book.sheetnames else None
    config = book["S3_Config"] if "S3_Config" in book.sheetnames else None
    refs = _config_refs(config)
    series = _series(days, tokens)
    nights = _night_pair(weekly)
    stages = _stage_pair(weekly)
    buckets = _activity_buckets(weekly)
    profile_map = _profile_map(profile)
    profile_view = _profile_view(profile_map)

    sections = []
    omitted = []
    for spec in rules["sections"]:
        page = str(spec["page"])
        layout = spec.get("layout") or "bars"
        block = pages.get(page)
        title = (block or {}).get("title") or spec.get("title") or page
        core = list(spec.get("core_headers") or [])
        core_missing = [name for name in core if _column_all_empty(days, name, tokens)]
        if core and len(core_missing) == len(core):
            omitted.append({"page": page, "title": title, "reason": "핵심 항목 없음: " + ", ".join(core_missing)})
            continue
        text_optional = layout in {"compare"}
        if not text_optional and (not block or (is_empty(block["stat"], tokens) and is_empty(block["meaning"], tokens))):
            omitted.append({"page": page, "title": title, "reason": "계산 결과가 비어 있어 생략"})
            continue
        block = block or {"stat": "", "meaning": "", "suggestion": ""}
        partial_names = [name for name in core if _column_some_empty(days, name, tokens)]
        status = "partial" if partial_names else "full"
        reason = (
            "일부 날짜만 있어 있는 데이터로 계산했습니다."
            if status == "partial"
            else "필요 항목이 있어 전체를 표시했습니다."
        )
        chart_header = _resolve_header(daily_headers, spec.get("chart_header") or "")
        chart = _chart(days, chart_header, spec.get("unit") or "", tokens, spec.get("ref") or "", spec.get("tone") or "", refs)
        sections.append(
            {
                "page": page,
                "layout": layout,
                "kicker": spec.get("kicker") or f"{page} · {title}",
                "question": spec.get("question") or title,
                "title": title,
                "status": status,
                "reason": reason,
                "unit": spec.get("unit") or "",
                "ref": spec.get("ref") or "",
                "stat": block["stat"] or ("" if text_optional else missing_text),
                "headline": _headline(block["stat"]),
                "meaning": block["meaning"] or ("" if text_optional else missing_text),
                "suggestion": block["suggestion"] or ("" if text_optional else missing_text),
                "cards": _routine_cards(block["suggestion"]) if layout == "routine" else [],
                "chart": chart,
            }
        )

    name = profile_view["name"]
    dates = [as_text(day["raw"].get("날짜"))[:10] for day in days if not is_empty(day["raw"].get("날짜"), tokens)]
    period = _period(dates)
    nights = _enrich_nights(nights, series)
    _attach_charts(sections, series)
    diagnosis = pages.get("16") or {}
    cover_line = _first_sentence(diagnosis.get("meaning") or "") or diagnosis.get("stat") or ""

    return {
        "document": rules["document"],
        "version": str(rules["version"]),
        "updated_at": str(rules["updated_at"]),
        "name": name or missing_text,
        "name_missing": not name,
        "period": period or missing_text,
        "profile": profile_view,
        "cover_line": cover_line,
        "series": series,
        "type_counts": _type_counts(series),
        "nights": nights,
        "stages": stages,
        "buckets": buckets,
        "refs": refs,
        "reference_rows": _reference_rows(series, refs),
        "signals": [row for row in series if row.get("signal")],
        "averages": _averages(series),
        "sections": sections,
        "omitted": omitted,
        "filled_count": input_filled,
        "total_count": input_total,
        "monitored_count": sum(monitored_flags),
        "day_count": len(days),
        "missing_labels": missing_labels,
        "generated_at": datetime.now(KST).strftime("%Y-%m-%d %H:%M"),
        "source_name": Path(workbook_path).name,
    }


def _require_sheets(book, names: list[str]) -> None:
    missing = [name for name in names if name not in book.sheetnames]
    if missing:
        raise ReportError(
            "수면 로직 엑셀 형식이 아닙니다. 필요한 시트가 없습니다: " + ", ".join(missing)
        )


def _headers(sheet, row: int) -> dict[str, int]:
    found = {}
    for col in range(1, sheet.max_column + 1):
        value = sheet.cell(row, col).value
        if value is None:
            continue
        text = str(value).strip()
        if text and not text.startswith("="):
            found[text] = col
    return found


def _data_rows(sheet) -> list[int]:
    rows = []
    for row in range(5, sheet.max_row + 1):
        value = sheet.cell(row, 1).value
        if is_empty(value):
            continue
        text = str(value).strip()
        if text.startswith(("※", "S1", "=")):
            continue
        rows.append(row)
    return rows


def _matching_daily_row(sheet, label: str, fallback: int) -> int:
    for row in range(5, sheet.max_row + 1):
        if as_text(sheet.cell(row, 1).value) == label:
            return row
    return fallback


def _profile_fields(sheet, tokens: set[str]) -> list[dict]:
    fields = []
    for row in range(5, sheet.max_row + 1):
        label = sheet.cell(row, 1).value
        if is_empty(label):
            continue
        text = str(label).strip()
        if text.startswith(("※", "=")) or "자동계산" in text:
            continue
        fields.append({"label": text, "value": "" if is_empty(sheet.cell(row, 2).value, tokens) else as_text(sheet.cell(row, 2).value)})
    return fields


def _count_inputs(days, raw_headers, profile_fields, tokens):
    skip = {"요일", "날짜"}
    headers = [name for name in raw_headers if name not in skip]
    total = 0
    filled = 0
    missing = []
    for day in days:
        for name in headers:
            total += 1
            if is_empty(day["raw"].get(name), tokens):
                missing.append(f"{day['label']} {name}")
            else:
                filled += 1
    for field in profile_fields:
        total += 1
        if field["value"]:
            filled += 1
        else:
            missing.append(field["label"])
    return total, filled, missing


def _column_all_empty(days, header: str, tokens: set[str]) -> bool:
    if not days:
        return True
    return all(is_empty(day["raw"].get(header), tokens) for day in days)


def _column_some_empty(days, header: str, tokens: set[str]) -> bool:
    values = [is_empty(day["raw"].get(header), tokens) for day in days]
    return any(values) and not all(values)


def _summary_pages(sheet) -> dict[str, dict]:
    pages = {}
    for row in range(5, sheet.max_row + 1):
        page = sheet.cell(row, 1).value
        if is_empty(page):
            continue
        pages[str(page).strip()] = {
            "title": as_text(sheet.cell(row, 2).value),
            "stat": as_text(sheet.cell(row, 3).value),
            "meaning": as_text(sheet.cell(row, 4).value),
            "suggestion": as_text(sheet.cell(row, 5).value),
        }
    return pages


def _resolve_header(headers: dict[str, int], wanted: str) -> str:
    if not wanted:
        return ""
    if wanted in headers:
        return wanted
    prefix = wanted.split("(")[0]
    for name in headers:
        if name.startswith(prefix):
            return name
    return ""


def _as_float(value):
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _pick_number(mapping: dict, name: str, tokens: set[str]):
    if name in mapping and not is_empty(mapping[name], tokens):
        return _as_float(mapping[name])
    for key, value in mapping.items():
        if str(key).startswith(name) and not is_empty(value, tokens):
            return _as_float(value)
    return None


def _series(days, tokens: set[str]) -> list[dict]:
    rows = []
    for day in days:
        if not day["monitored"]:
            continue
        raw, daily = day["raw"], day["daily"]
        deep = _pick_number(raw, "깊은수면_분", tokens)
        light = _pick_number(raw, "얕은수면_분", tokens)
        rem = _pick_number(raw, "렘수면_분", tokens)
        stage_total = (deep or 0) + (light or 0) + (rem or 0)
        signal = daily.get("이상신호_해당여부")
        flagged = signal is True or str(signal).strip() in {"True", "TRUE", "예", "Y", "1"}
        rows.append(
            {
                "label": day["label"],
                "tst": _pick_number(daily, "총수면시간_TST_분", tokens),
                "deep": deep,
                "light": light,
                "rem": rem,
                "deep_share": round((deep or 0) / stage_total * 100, 1) if stage_total else 0,
                "light_share": round((light or 0) / stage_total * 100, 1) if stage_total else 0,
                "rem_share": round((rem or 0) / stage_total * 100, 1) if stage_total else 0,
                "se": _pick_number(daily, "수면효율_최종_%", tokens),
                "sol": _pick_number(daily, "입면잠복기_SOL_분", tokens),
                "awake": _pick_number(daily, "각성횟수", tokens),
                "awake_min": _pick_number(daily, "평균각성지속시간_분", tokens),
                "hr": _pick_number(daily, "심박수_HR", tokens),
                "rr": _pick_number(daily, "호흡수_RR", tokens),
                "spo2": _pick_number(daily, "SpO2_평균", tokens),
                "bed": as_text(raw.get("취침시각")),
                "wake": as_text(raw.get("기상시각")),
                "bed_min": _pick_number(daily, "취침시각_분환산", tokens),
                "wake_min": _pick_number(daily, "기상시각_분", tokens),
                "activity": _pick_number(daily, "활동시간_분", tokens),
                "kcal": _pick_number(daily, "소모칼로리_기기값", tokens),
                "vitality": _pick_number(daily, "활력점수_기기값", tokens),
                "tib": _pick_number(daily, "침상시간_TIB_분", tokens),
                "steps": _pick_number(raw, "걸음수", tokens),
                "date": _month_day(as_text(raw.get("날짜"))),
                "sleep_type": as_text(daily.get("수면유형_분류")),
                "signal": flagged,
            }
        )
    if not rows:
        return []
    peak = max((row["tst"] or 0) for row in rows) or 1
    span_start, span_end = -180, 660
    span = span_end - span_start
    for row in rows:
        row["tst_pct"] = round((row["tst"] or 0) / peak * 100, 1)
        bed_min, wake_min = row["bed_min"], row["wake_min"]
        if bed_min is None or wake_min is None:
            row["clock_left"] = 0
            row["clock_width"] = 0
            continue
        left = (bed_min - span_start) / span * 100
        width = (wake_min - bed_min) / span * 100
        row["clock_left"] = round(max(0, min(left, 100)), 1)
        row["clock_width"] = round(max(0, min(width, 100 - max(left, 0))), 1)
    return rows


def _profile_map(sheet) -> dict[str, str]:
    found = {}
    for row in range(5, sheet.max_row + 1):
        label = as_text(sheet.cell(row, 1).value)
        if not label or label.startswith(("※", "=")):
            continue
        found[label] = as_text(sheet.cell(row, 2).value)
    return found


def _label_value(mapping: dict[str, str], prefix: str) -> str:
    for key, value in mapping.items():
        if key.startswith(prefix) and value:
            return value
    return ""


def _profile_view(mapping: dict[str, str]) -> dict:
    bmi = _label_value(mapping, "BMI")
    bmr = _label_value(mapping, "기초대사량")
    try:
        bmi = f"{float(bmi):.1f}"
    except ValueError:
        pass
    try:
        bmr = f"{round(float(bmr)):,}"
    except ValueError:
        pass
    age = _label_value(mapping, "나이")
    weight = _label_value(mapping, "체중")
    height = _label_value(mapping, "키")
    return {
        "name": _label_value(mapping, "이름"),
        "sex": _label_value(mapping, "성별"),
        "age": age.split(".")[0] if age else "",
        "weight": weight.split(".")[0] if weight.endswith(".0") else weight,
        "height": height.split(".")[0] if height.endswith(".0") else height,
        "bmi": bmi,
        "bmr": bmr,
    }


def _config_refs(sheet) -> dict:
    wanted = {
        "se": "수면효율 (SE)",
        "sol": "입면잠복기 (SOL)",
        "awake": "각성횟수",
        "hr": "심박수 (HR",
        "rr": "호흡수 (RR",
        "spo2": "SpO2",
        "activity": "활동시간 목표",
    }
    found = {key: {"low": None, "high": None, "target": None} for key in wanted}
    if sheet is None:
        return found
    for row in range(5, sheet.max_row + 1):
        label = as_text(sheet.cell(row, 1).value)
        for key, prefix in wanted.items():
            if label.startswith(prefix):
                found[key] = {
                    "low": _as_float(sheet.cell(row, 2).value),
                    "high": _as_float(sheet.cell(row, 3).value),
                    "target": _as_float(sheet.cell(row, 4).value),
                }
    return found


def _sheet_row(sheet, label: str):
    if sheet is None:
        return None
    for row in range(1, sheet.max_row + 1):
        if as_text(sheet.cell(row, 1).value) == label:
            return row
    return None


def _night_pair(sheet) -> dict:
    good = _sheet_row(sheet, "좋았던 밤 (SE 최댓값)")
    bad = _sheet_row(sheet, "나빴던 밤 (SE 최솟값)")

    def pack(row):
        if row is None:
            return {}
        return {
            "day": as_text(sheet.cell(row, 2).value),
            "se": as_text(sheet.cell(row, 3).value),
            "bed": as_text(sheet.cell(row, 4).value),
            "wake": as_text(sheet.cell(row, 5).value),
        }

    return {"good": pack(good), "bad": pack(bad)}


def _stage_pair(sheet) -> dict:
    long_row = _sheet_row(sheet, "TST 최댓값 요일")
    short_row = _sheet_row(sheet, "TST 최솟값 요일")

    def pack(row):
        if row is None:
            return {}
        return {
            "day": as_text(sheet.cell(row, 2).value),
            "deep": as_text(sheet.cell(row, 3).value),
            "light": as_text(sheet.cell(row, 4).value),
            "rem": as_text(sheet.cell(row, 5).value),
        }

    return {"long": pack(long_row), "short": pack(short_row)}


def _activity_buckets(sheet) -> list[dict]:
    if sheet is None:
        return []
    rows = []
    for label in ("저활동", "중활동", "고활동"):
        row = _sheet_row(sheet, label)
        if row is None:
            continue
        rows.append(
            {
                "name": label,
                "days": as_text(sheet.cell(row, 2).value),
                "se": as_text(sheet.cell(row, 3).value),
                "tst": as_text(sheet.cell(row, 4).value),
            }
        )
    return rows


def _month_day(text: str) -> str:
    parts = text[:10].split("-")
    if len(parts) != 3 or not parts[0].isdigit():
        return text
    return f"{int(parts[1])}/{int(parts[2])}"


def _period(dates: list[str]) -> str:
    if not dates:
        return ""
    start = dates[0].replace("-", ".")
    if len(dates) == 1:
        return start
    end = dates[-1]
    if start[:4] == end[:4]:
        return f"{start} – {end[5:7]}.{end[8:10]}"
    return f"{start} – {end.replace('-', '.')}"


def _enrich_nights(nights: dict, series: list[dict]) -> dict:
    for key in ("good", "bad"):
        night = nights.get(key) or {}
        match = next((row for row in series if row["label"] == night.get("day")), None)
        if not match:
            continue
        night["date"] = match.get("date") or ""
        night["tst"] = hours_label(match.get("tst"))
        night["sol"] = _plain(match.get("sol"))
        night["deep"] = hours_label(match.get("deep"))
        night["light"] = hours_label(match.get("light"))
        night["rem"] = hours_label(match.get("rem"))
        night["bed"] = night.get("bed") or match.get("bed") or ""
        night["wake"] = night.get("wake") or match.get("wake") or ""
    return nights


def _plain(value) -> str:
    if value is None:
        return ""
    number = float(value)
    if number.is_integer():
        return str(int(number))
    return f"{number:.1f}"


def _averages(series: list[dict]) -> dict:
    def average(key: str):
        values = [float(row[key]) for row in series if row.get(key) is not None]
        if not values:
            return None
        return sum(values) / len(values)

    tib = average("tib")
    steps = average("steps")
    awake = average("awake")
    return {
        "tib": hours_label(tib) if tib is not None else "",
        "steps": f"{round(steps):,}" if steps is not None else "",
        "awake": f"{awake:.1f}" if awake is not None else "",
    }


def _reference_rows(series: list[dict], refs: dict) -> list[dict]:
    def average(key: str):
        values = [float(row[key]) for row in series if row.get(key) is not None]
        if not values:
            return None
        return sum(values) / len(values)

    specs = [
        ("수면효율", "se", "%", "min"),
        ("입면잠복기", "sol", "분", "max"),
        ("각성횟수", "awake", "회", "max"),
        ("수면 중 심박수", "hr", "bpm", "band"),
        ("호흡수", "rr", "회/분", "band"),
        ("산소포화도", "spo2", "%", "min"),
    ]
    rows = []
    for label, key, unit, mode in specs:
        value = average(key)
        if value is None:
            continue
        rule = refs.get(key) or {}
        low, high = rule.get("low"), rule.get("high")
        shown = _plain(value) + unit
        guide = ""
        gap = "범위 안"
        tone = "ok"
        if mode == "min" and low is not None:
            guide = f"{_plain(low)}{unit} 이상"
            delta = value - low
            if delta < -0.05:
                gap = f"{_plain(abs(delta))}{unit} 낮음"
                tone = "low"
        elif mode == "max" and high is not None:
            guide = f"{_plain(high)}{unit} 이내"
            delta = value - high
            if delta > 0.05:
                gap = f"{_plain(delta)}{unit} 김" if unit == "분" else f"{_plain(delta)}{unit} 초과"
                tone = "low"
        elif mode == "band":
            guide = f"{_plain(low) if low is not None else ''}–{_plain(high) if high is not None else ''}{unit}"
            if low is not None and value < low:
                gap = "범위 아래"
                tone = "low"
            elif high is not None and value > high:
                gap = "범위 위"
                tone = "low"
        rows.append({"label": label, "value": shown, "guide": guide, "gap": gap, "tone": tone})
    return rows


def _attach_charts(sections: list[dict], series: list[dict]) -> None:
    palette = {
        "se": ("#a0c8c0", "#184048"),
        "sol": ("#a0c8c0", "#184048"),
        "awake": ("#a0c8c0", "#184048"),
        "spo2": ("#e0c0a8", "#8a5848"),
        "hr": ("#e0c0a8", "#8a5848"),
        "rr": ("#d7c3ae", "#8a5848"),
        "vitality": ("#e0c0a8", "#8a5848"),
    }
    keys = {
        "04": "se",
        "05": "sol",
        "06": "awake",
        "08": "hr",
        "09": "rr",
        "10": "spo2",
        "15": "vitality",
    }
    for section in sections:
        page = section["page"]
        if section["layout"] == "stacked":
            totals = [row["tst"] for row in series if row.get("tst")]
            section["hero"] = hours_label(sum(totals) / len(totals)) if totals else section.get("headline")
            section["svg"] = stacked_svg(series)
        elif page in keys:
            key = keys[page]
            bar, line = palette[key]
            section["svg"] = combo_svg(
                [{"label": row["label"], "value": row.get(key)} for row in series],
                bar=bar,
                line=line,
            )
            if key == "vitality":
                totals = [row["vitality"] for row in series if row.get("vitality") is not None]
                if len(totals) >= 2:
                    section["hero"] = f"{_plain(totals[0])}점 → {_plain(totals[-1])}점"
        elif section["layout"] == "quad":
            section["svg_tib"] = dual_svg(series, bar_key="tib", line_key="se", bar_fmt=hours_label, line_fmt=lambda value: f"{value:.0f}%")
            section["svg_activity"] = dual_svg(series, bar_key="activity", line_key="sol", line="#c47a4a", bar_fmt=lambda value: f"{value:.0f}", line_fmt=lambda value: f"{value:.0f}분")
        elif section["layout"] == "clock":
            section["svg_bed"] = line_svg([{"label": row["label"], "value": row.get("bed_min")} for row in series])
            section["svg_wake"] = line_svg(
                [{"label": row["label"], "value": row.get("wake_min")} for row in series],
                color="#c47a4a",
            )
        elif section["layout"] == "activity":
            section["svg"] = dual_svg(
                series,
                bar_key="activity",
                line_key="steps",
                line="#8a5848",
                bar_fmt=lambda value: f"{value:.0f}분",
                line_fmt=lambda value: f"{value:,.0f}",
            )
        elif section["layout"] == "signals":
            flagged = next((row["label"] for row in series if row.get("signal")), "")
            section["svg_hr"] = combo_svg([{"label": row["label"], "value": row.get("hr")} for row in series], bar="#d5e4e0", line="#184048", highlight=flagged, height=120)
            section["svg_rr"] = combo_svg([{"label": row["label"], "value": row.get("rr")} for row in series], bar="#d5e4e0", line="#184048", highlight=flagged, height=120)
            section["svg_spo2"] = combo_svg([{"label": row["label"], "value": row.get("spo2")} for row in series], bar="#d5e4e0", line="#184048", highlight=flagged, height=120)


def _type_counts(series: list[dict]) -> list[dict]:
    counts: dict[str, int] = {}
    for row in series:
        name = row.get("sleep_type") or ""
        if not name:
            continue
        counts[name] = counts.get(name, 0) + 1
    return [{"name": name, "days": days} for name, days in counts.items()]


def _headline(stat: str) -> str:
    text = (stat or "").strip()
    if not text:
        return ""
    first = text.split("·")[0].strip()
    for prefix in ("평균 ", "1회당 ", "이상신호 "):
        if first.startswith(prefix):
            return first[len(prefix):].strip()
    return first


def _first_sentence(text: str) -> str:
    if not text:
        return ""
    for mark in ("습니다.", "있습니다.", "않았습니다.", "않았습니다"):
        index = text.find(mark)
        if index != -1:
            return text[: index + len(mark)]
    return text


def _routine_cards(text: str) -> list[dict]:
    if not text:
        return []
    marks = "①②③④"
    if not any(mark in text for mark in marks):
        return [{"mark": "", "text": text}]
    chunks = []
    current = ""
    mark = ""
    for char in text:
        if char in marks:
            if current.strip():
                chunks.append({"mark": mark, "text": current.strip()})
            mark = char
            current = ""
        else:
            current += char
    if current.strip():
        chunks.append({"mark": mark, "text": current.strip()})
    return chunks


def _tone(value: float, ref_name: str, mode: str, refs: dict) -> str:
    spec = refs.get(ref_name) or {}
    low, high = spec.get("low"), spec.get("high")
    if mode == "min" and low is not None:
        return "ok" if value >= low else "low"
    if mode == "max" and high is not None:
        return "ok" if value <= high else "low"
    if mode == "band":
        if low is not None and value < low:
            return "low"
        if high is not None and value > high:
            return "low"
        if low is not None or high is not None:
            return "ok"
    return "neutral"


def _chart(days, header: str, unit: str, tokens: set[str], ref_name: str = "", tone_mode: str = "", refs: dict | None = None):
    if not header:
        return None
    points = []
    for day in days:
        if not day["monitored"] and header not in day["daily"]:
            continue
        value = day["daily"].get(header)
        if is_empty(value, tokens) or isinstance(value, str):
            continue
        try:
            number = float(value)
        except (TypeError, ValueError):
            continue
        points.append(
            {
                "label": day["label"],
                "value": number,
                "display": _format_chart(number, unit),
                "tone": _tone(number, ref_name, tone_mode, refs or {}),
            }
        )
    if not points:
        return None
    peak = max(abs(point["value"]) for point in points) or 1
    for point in points:
        if unit == "%":
            point["pct"] = round(min(abs(point["value"]), 100), 1)
        else:
            point["pct"] = round(abs(point["value"]) / peak * 100, 1)
    ref = (refs or {}).get(ref_name) or {}
    marker = ref.get("low") if tone_mode == "min" else ref.get("high") if tone_mode == "max" else None
    return {"title": header, "unit": unit, "marker": marker if unit == "%" else None, "points": points}


def _format_chart(number: float, unit: str) -> str:
    if unit == "%":
        return f"{number:.0f}%"
    if number.is_integer():
        text = f"{int(number):,}"
    else:
        text = f"{number:.1f}"
    return f"{text}{unit}" if unit else text
