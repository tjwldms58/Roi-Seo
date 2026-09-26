import io
import shutil
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from sleep_report.xiaomi import fill_logic, list_people

from recovery_report.engine import LOGIC_PATH as RECOVERY_LOGIC
from recovery_report.engine import build_model as build_recovery
from recovery_report.render import render_html as recovery_html
from recovery_report.render import render_pdf as recovery_pdf
from recovery_report.web import create_app as recovery_app
from sleep_report.engine import LOGIC_PATH as SLEEP_LOGIC
from sleep_report.engine import build_model as build_sleep
from sleep_report.render import render_html as sleep_html
from sleep_report.render import render_pdf as sleep_pdf
from sleep_report.web import create_app as sleep_app


def _copy(src: Path, folder: Path, name: str = "input.xlsx") -> Path:
    dest = folder / name
    shutil.copy(src, dest)
    return dest


def test_sleep_sample_uses_workbook_text(tmp_path):
    model = build_sleep(_copy(SLEEP_LOGIC, tmp_path))
    pages = {item["page"]: item for item in model["sections"]}
    assert "03" in pages
    assert "301" in pages["03"]["meaning"] or "301" in pages["03"]["stat"]
    assert "06+" in pages
    assert model["filled_count"] > 0
    assert model["total_count"] >= model["filled_count"]
    assert model["version"] == "19차"
    html = sleep_html(model)
    assert "healmarucare" in html
    assert "NotoSansKR-Regular.woff2" in html
    assert ".otf" not in html
    assert "데이터 기반 참고 리포트" in html
    assert "생성 정보" not in html
    assert "width: 210mm" in html and "height: 297mm" in html
    pdf = sleep_pdf(html)
    assert pdf.startswith(b"%PDF")


def test_sleep_omits_wake_duration_without_measured_time(tmp_path):
    path = _copy(SLEEP_LOGIC, tmp_path)
    book = openpyxl.load_workbook(path)
    sheet = book["S1_Input_Raw"]
    for row in range(5, 12):
        sheet.cell(row, 18).value = None
    book.save(path)
    model = build_sleep(path)
    omitted = {item["page"] for item in model["omitted"]}
    assert "06+" in omitted
    assert "03" in {item["page"] for item in model["sections"]}


def test_sleep_reference_comes_from_config_sheet(tmp_path):
    path = _copy(SLEEP_LOGIC, tmp_path)
    book = openpyxl.load_workbook(path)
    config = book["S3_Config"]
    config["B6"].value = 50
    config["D6"].value = 50
    book.save(path)
    model = build_sleep(path)
    page = next(item for item in model["sections"] if item["page"] == "04")
    assert "50%" in page["meaning"] or "50%" in page["stat"]


def test_sleep_blank_core_omits_sections(tmp_path):
    path = _copy(SLEEP_LOGIC, tmp_path)
    book = openpyxl.load_workbook(path)
    sheet = book["S1_Input_Raw"]
    for row in range(5, 12):
        for col in range(3, 20):
            sheet.cell(row, col).value = None
    book.save(path)
    model = build_sleep(path)
    assert model["sections"] == []
    assert model["omitted"]
    assert model["filled_count"] < model["total_count"]


def test_recovery_sample_keeps_capability_page(tmp_path):
    model = build_recovery(_copy(RECOVERY_LOGIC, tmp_path))
    ids = [item["id"] for item in model["sections"]]
    assert "cover" in ids
    assert "capability" in ids
    assert "cancer" in ids
    assert model["content"]["name"].startswith("김OO")
    assert model["content"]["arm_delta"].startswith("-2")
    # 양식의 예비 입력칸(추가 지표 행)은 샘플에서 비어 있다. 섹션은 그대로 그린다.
    assert 0 < model["filled_count"] < model["total_count"]
    assert model["omitted"] == []
    html = recovery_html(model)
    assert "HEALMARU CARE" in html
    assert "데이터 기반 참고 리포트" in html
    assert "생성 정보" not in html
    assert "이런 변화가 있었어요" in html
    assert "서지은 간호사" in html
    assert "width: 210mm" in html and "height: 297mm" in html
    pdf = recovery_pdf(html)
    assert pdf.startswith(b"%PDF")


def test_recovery_omits_capability_when_names_missing(tmp_path):
    path = _copy(RECOVERY_LOGIC, tmp_path)
    book = openpyxl.load_workbook(path)
    sheet = book["04_운동별체력변화"]
    for row in range(7, 13):
        for col in range(1, 5):
            sheet.cell(row, col).value = None
    book.save(path)
    model = build_recovery(path)
    assert "capability" not in [item["id"] for item in model["sections"]]
    assert any(item["id"] == "capability" for item in model["omitted"])
    assert "summary" in [item["id"] for item in model["sections"]]
    assert model["filled_count"] < model["total_count"]


def test_one_site_separates_sleep_and_recovery():
    from site_app.web import create_site

    client = create_site().test_client()
    home = client.get("/").get_data(as_text=True)
    assert "수면 레포트 만들기" in home
    assert "동료와 함께 쓰기" in home
    assert "종합 레포트 만들기" in home
    font = client.get("/fonts/NotoSansKR-Regular.woff2")
    assert font.status_code == 200
    assert font.mimetype == "font/woff2"
    assert font.data[:4] == b"wOF2"
    sleep = client.get("/sleep/")
    assert sleep.status_code == 200
    sleep_html = sleep.get_data(as_text=True)
    assert "샤오미 수면 데이터" in sleep_html
    assert 'action="/sleep/generate"' in sleep_html
    recovery = client.get("/recovery/")
    assert recovery.status_code == 200
    recovery_html = recovery.get_data(as_text=True)
    assert "종합 레포트 양식" in recovery_html
    assert 'action="/recovery/generate"' in recovery_html


def test_web_uploads(tmp_path, monkeypatch):
    sleep_client = sleep_app().test_client()
    page = sleep_client.get("/")
    assert "19차" in page.get_data(as_text=True)
    with SLEEP_LOGIC.open("rb") as handle:
        response = sleep_client.post(
            "/generate",
            data={"file": (handle, "sleep.xlsx")},
            content_type="multipart/form-data",
            follow_redirects=True,
        )
    assert "PDF 받기" in response.get_data(as_text=True)

    recovery_client = recovery_app().test_client()
    with RECOVERY_LOGIC.open("rb") as handle:
        response = recovery_client.post(
            "/generate",
            data={"file": (handle, "recovery.xlsx")},
            content_type="multipart/form-data",
            follow_redirects=True,
        )
    assert "PDF 받기" in response.get_data(as_text=True)
    template = recovery_client.get("/template")
    assert template.status_code == 200
    assert template.data[:2] == b"PK"


def _xiaomi_book(rows: list[tuple]) -> bytes:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "샤오미데이터"
    headers = [
        "날짜",
        "고객명",
        "나이",
        "성별",
        "잠든 시각",
        "깬 시각",
        "수면준비시간 (분)",
        "깊은수면 시간",
        "얕은수면 시간",
        "REM 시간",
        "각성 (분)",
        "수면효율성 (%)",
        "수면시 평균 심박수 (BPM)",
        "평균 호흡수 (BPM)",
        "평균 혈중 산소포화도 (%)",
        "활력점수",
        "체중(KG)",
        "키(CM)",
    ]
    sheet.append(headers)
    for row in rows:
        sheet.append(list(row))
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_xiaomi_table_fills_logic_inputs_without_sample_profile(tmp_path):
    source = tmp_path / "xiaomi.xlsx"
    source.write_bytes(
        _xiaomi_book(
            [
                ("2026-09-08", "테스트갑", 65, "여", "01:48", "07:06", 14, "0h 36m", "3h 17m", "1h 11m", 14, 91, 67, 15, 96, 21, None, None),
                ("2026-09-09", "테스트을", 52, "여", "23:03", "06:08", 17, "1h 12m", "3h 53m", "1h 50m", 10, 50, 66, 18, 96, 16, 60, 160),
            ]
        )
    )
    people = list_people(source)
    assert [item["name"] for item in people] == ["테스트갑", "테스트을"]
    assert people[0]["nights"] == 1
    dest = tmp_path / "filled.xlsx"
    fill_logic(source, dest, "테스트갑")
    book = openpyxl.load_workbook(dest)
    raw = book["S1_Input_Raw"]
    headers = {raw.cell(4, col).value: col for col in range(1, 20)}
    assert raw.cell(5, headers["요일"]).value == "화"
    assert raw.cell(5, headers["날짜"]).value.date() == date(2026, 9, 8)
    assert raw.cell(5, headers["취침시각"]).value.hour == 1
    assert raw.cell(5, headers["깊은수면_분"]).value == 36
    assert raw.cell(5, headers["렘수면_분"]).value == 71
    assert raw.cell(5, headers["각성시간_실측_분(기기값,선택)"]).value == 14
    assert raw.cell(5, headers["각성횟수_회"]).value is None
    assert raw.cell(6, headers["요일"]).value is None
    profile = book["S2_Input_Profile"]
    assert profile["B5"].value == "테스트갑"
    assert profile["B7"].value == 65
    assert profile["B8"].value is None
    assert profile["B11"].value is None
    other = tmp_path / "other.xlsx"
    fill_logic(source, other, "테스트을")
    kept = openpyxl.load_workbook(other)["S2_Input_Profile"]
    assert kept["B8"].value == 60
    assert kept["B9"].value == 160
    assert str(kept["B11"].value).startswith("=")


def test_xiaomi_keeps_latest_seven_days(tmp_path):
    rows = [
        (f"2026-09-{day:02d}", "테스트갑", 65, "여", "01:00", "07:00", 10, "0h 30m", "3h 0m", "1h 0m", 5, 80, 60, 15, 96, 20, None, None)
        for day in range(8, 16)
    ]
    source = tmp_path / "xiaomi.xlsx"
    source.write_bytes(_xiaomi_book(rows))
    dest = tmp_path / "filled.xlsx"
    notice = fill_logic(source, dest, "테스트갑")
    assert "최근 7일" in notice
    book = openpyxl.load_workbook(dest)
    raw = book["S1_Input_Raw"]
    assert raw.cell(5, 2).value.date() == date(2026, 9, 9)
    assert raw.cell(11, 2).value.date() == date(2026, 9, 15)


def test_xiaomi_upload_asks_for_a_person_then_builds_report():
    payload = _xiaomi_book(
        [
            ("2026-09-16", "테스트갑", 65, "여", "01:48", "07:06", 14, "0h 36m", "3h 17m", "1h 11m", 14, 91, 67, 15, 96, 21, None, None),
            ("2026-09-16", "테스트을", 52, "여", "23:03", "06:08", 17, "1h 12m", "3h 53m", "1h 50m", 10, 50, 66, 18, 96, 16, None, None),
        ]
    )
    client = sleep_app().test_client()
    response = client.post(
        "/generate",
        data={"file": (io.BytesIO(payload), "샤오미.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    page = response.get_data(as_text=True)
    assert "테스트갑" in page
    assert "테스트을" in page
    assert "PDF 받기" not in page
    chosen = client.post(
        response.request.path,
        data={"person": "0"},
        follow_redirects=True,
    )
    result = chosen.get_data(as_text=True)
    assert "PDF 받기" in result
    assert "측정 1/1일" in result
    html_path = chosen.request.path.replace("/result/", "/download/") + "/html"
    report = client.get(html_path).get_data(as_text=True)
    assert "테스트갑" in report
    assert "예시 대상자" not in report
    assert "91%" in report
